#!/usr/bin/env python3
"""Export the RoadOcc framework as an editable PowerPoint slide.

The slide mirrors the publication SVG but keeps text, boxes, arrows, and
diagram glyphs as native PowerPoint objects. Only camera/model diagnostics are
embedded as raster image layers because they are experimental visual evidence.
"""

from __future__ import annotations

import argparse
import io
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import draw_roadocc_framework as source  # noqa: E402


CANVAS_W = 1800.0
CANVAS_H = 525.0
SLIDE_W_IN = 13.333333
SLIDE_H_IN = SLIDE_W_IN * CANVAS_H / CANVAS_W
FONT = "Times New Roman"


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor.from_string(value.upper())


def x(value: float):
    return Inches(value / CANVAS_W * SLIDE_W_IN)


def y(value: float):
    return Inches(value / CANVAS_H * SLIDE_H_IN)


def pt(value: float) -> Pt:
    return Pt(value * SLIDE_W_IN * 72.0 / CANVAS_W)


def set_shape_name(shape, name: str | None) -> None:
    if name:
        shape.name = name


def set_line_dash(line, dash: bool) -> None:
    if dash:
        line.dash_style = MSO_LINE_DASH_STYLE.DASH


def add_arrowhead(line, arrow: bool) -> None:
    if not arrow:
        return
    line_element = line._get_or_add_ln()
    tail = line_element.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    if tail is None:
        tail = OxmlElement("a:tailEnd")
        line_element.append(tail)
    tail.set("type", "triangle")
    tail.set("w", "sm")
    tail.set("len", "sm")


def add_rect(slide, left: float, top: float, width: float, height: float,
             *, fill: str = "#FFFFFF", stroke: str = "#A8B3C2",
             sw: float = 1.3, radius: bool = False, dash: bool = False,
             name: str | None = None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x(left), y(top), x(width), y(height))
    set_shape_name(shape, name)
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(stroke)
    shape.line.width = max(Pt(0.7), pt(sw))
    set_line_dash(shape.line, dash)
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except (IndexError, ValueError):
            pass
    return shape


def add_circle(slide, cx: float, cy: float, radius: float, *, fill: str,
               stroke: str = "#FFFFFF", sw: float = 1.0,
               name: str | None = None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, x(cx - radius), y(cy - radius),
        x(2 * radius), y(2 * radius))
    set_shape_name(shape, name)
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(stroke)
    shape.line.width = max(Pt(0.6), pt(sw))
    return shape


def add_text(slide, cx: float, baseline: float, value: str, *,
             size: float = 18, color: str = source.INK,
             bold: bool = False, italic: bool = False,
             align: str = "center", width: float = 220,
             height: float | None = None, name: str | None = None):
    if height is None:
        height = size * 1.35
    if align == "center":
        left = cx - width / 2
        paragraph_alignment = PP_ALIGN.CENTER
    elif align == "left":
        left = cx
        paragraph_alignment = PP_ALIGN.LEFT
    else:
        left = cx - width
        paragraph_alignment = PP_ALIGN.RIGHT
    top = baseline - size * 1.02
    box = slide.shapes.add_textbox(x(left), y(top), x(width), y(height))
    set_shape_name(box, name)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = paragraph_alignment
    paragraph.space_before = paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return box


def add_multiline(slide, cx: float, baseline: float, lines: list[str], *,
                  size: float = 17, color: str = source.INK,
                  bold: bool = False, italic: bool = False,
                  width: float = 220, line_height: float = 19,
                  name: str | None = None):
    height = line_height * len(lines) + 8
    box = slide.shapes.add_textbox(
        x(cx - width / 2), y(baseline - size), x(width), y(height))
    set_shape_name(box, name)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, value in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = pt(line_height)
        run = paragraph.add_run()
        run.text = value
        run.font.name = FONT
        run.font.size = pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)
    return box


def add_line(slide, x1: float, y1: float, x2: float, y2: float, *,
             color: str = source.INK, sw: float = 1.8,
             arrow: bool = False, dash: bool = False,
             name: str | None = None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, x(x1), y(y1), x(x2), y(y2))
    set_shape_name(shape, name)
    shape.line.color.rgb = rgb(color)
    shape.line.width = max(Pt(0.75), pt(sw))
    set_line_dash(shape.line, dash)
    add_arrowhead(shape.line, arrow)
    return shape


def add_elbow(slide, points: list[tuple[float, float]], *,
              color: str = source.INK, sw: float = 1.8,
              arrow: bool = False, dash: bool = False,
              name: str | None = None):
    shapes = []
    for index, (start, end) in enumerate(zip(points, points[1:])):
        shapes.append(add_line(
            slide, start[0], start[1], end[0], end[1], color=color,
            sw=sw, arrow=arrow and index == len(points) - 2, dash=dash,
            name=f"{name}-{index + 1}" if name else None))
    return shapes


def add_picture(slide, left: float, top: float, width: float, height: float,
                image: Image.Image, *, stroke: str = "#D7DEE8",
                name: str | None = None):
    stream = io.BytesIO()
    image.save(stream, format="PNG", optimize=True)
    stream.seek(0)
    picture = slide.shapes.add_picture(
        stream, x(left), y(top), width=x(width), height=y(height))
    set_shape_name(picture, name)
    image_ratio = image.width / image.height
    frame_ratio = width / height
    if image_ratio > frame_ratio:
        visible = frame_ratio / image_ratio
        picture.crop_left = picture.crop_right = (1.0 - visible) / 2.0
    else:
        visible = image_ratio / frame_ratio
        picture.crop_top = picture.crop_bottom = (1.0 - visible) / 2.0
    border = add_rect(
        slide, left, top, width, height, fill="#FFFFFF", stroke=stroke,
        sw=0.7, radius=False, name=f"{name}-border" if name else None)
    border.fill.background()
    return picture


def add_module_panel(slide, left: float, top: float, width: float,
                     height: float, *, color: str, fill: str,
                     title: str, question: str, tab_width: float,
                     name: str):
    add_rect(slide, left, top, width, height, fill=fill, stroke=color,
             sw=1.7, radius=True, name=name)
    add_rect(slide, left + 10, top - 12, tab_width, 29, fill=color,
             stroke=color, sw=1.0, radius=True, name=f"{name}-tab")
    add_text(slide, left + 10 + tab_width / 2, top + 8, title,
             size=18, color="#FFFFFF", bold=True, width=tab_width - 8,
             name=f"{name}-title")
    add_text(slide, left + width - 10, top + 8, question, size=17,
             color=source.INK, bold=True, italic=True, align="right",
             width=150, name=f"{name}-question")


def add_voxel_cube(slide):
    add_rect(slide, 354, 148, 74, 74, fill="#A9C3E3",
             stroke="#536171", sw=1.3, name="current-voxel-front")
    top = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM, x(354), y(128), x(98), y(20))
    top.name = "current-voxel-top"
    top.shadow.inherit = False
    top.fill.solid()
    top.fill.fore_color.rgb = rgb("#D7E5F4")
    top.line.color.rgb = rgb("#536171")
    top.line.width = Pt(0.8)
    side = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM, x(428), y(128), x(24), y(94))
    side.name = "current-voxel-side"
    side.shadow.inherit = False
    side.fill.solid()
    side.fill.fore_color.rgb = rgb("#7698C2")
    side.line.color.rgb = rgb("#536171")
    side.line.width = Pt(0.8)
    for index in range(1, 4):
        add_line(slide, 354 + 74 * index / 4, 148,
                 354 + 74 * index / 4, 222, color="#FFFFFF", sw=0.7)
        add_line(slide, 354, 148 + 74 * index / 4,
                 428, 148 + 74 * index / 4, color="#FFFFFF", sw=0.7)


def build_slide(prs: Presentation, traces: dict[str, Image.Image]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb("#FFFFFF")

    # Input and camera-geometry evidence.
    add_text(slide, 118, 32, "Roadside images", size=24, bold=True,
             width=210, name="roadside-images-heading")
    add_rect(slide, 16, 50, 205, 215, fill="#FFFFFF", stroke=source.LINE,
             sw=1.5, radius=True, name="input-images")
    add_picture(slide, 28, 64, 181, 176, traces["cameras"],
                name="camera-trace")
    add_text(slide, 118, 257, "current frame  t", size=17,
             color=source.MUTED, width=180, name="current-frame-label")
    add_rect(slide, 34, 287, 260, 210, fill="#FFFFFF", stroke=source.LINE,
             sw=1.4, radius=True, name="camera-geometry")
    add_picture(slide, 46, 299, 236, 168, traces["camera_fov"],
                name="camera-fov-trace")
    add_text(slide, 164, 488, "camera geometry", size=17, bold=True,
             width=200, name="camera-geometry-label")
    add_elbow(slide, [(164, 287), (164, 276), (278, 276), (278, 216)],
              color=source.NEUTRAL, sw=1.5, arrow=True, dash=True,
              name="calibration-depth-path")
    add_text(slide, 222, 269, "calibration + depth", size=14.5,
             color=source.MUTED, bold=True, width=170,
             name="calibration-depth-label")

    # Encoder and RoadOcc boundary.
    add_line(slide, 222, 157, 247, 157, sw=2.3, arrow=True,
             name="input-to-encoder")
    encoder = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM, x(250), y(100), x(56), y(114))
    encoder.name = "image-encoder"
    encoder.shadow.inherit = False
    encoder.fill.solid()
    encoder.fill.fore_color.rgb = rgb(source.ENCODER_FILL)
    encoder.line.color.rgb = rgb(source.ENCODER)
    encoder.line.width = Pt(1.0)
    add_multiline(slide, 278, 144, ["image", "encoder"], size=17,
                  color=source.ENCODER, bold=True, width=60,
                  line_height=20, name="image-encoder-label")
    add_line(slide, 307, 157, 337, 157, sw=2.3, arrow=True,
             name="encoder-to-voxel")

    add_rect(slide, 330, 34, 1205, 485, fill="#FFFFFF", stroke="#64748B",
             sw=1.6, radius=True, name="roadocc-decoder")
    add_text(slide, 745, 58, "RoadOcc coarse-to-fine decoder", size=22,
             bold=True, width=430, name="decoder-heading")
    add_text(slide, 1110, 58, "shared across", size=16,
             color=source.MUTED, bold=True, width=110,
             name="shared-across-label")
    for index, label in enumerate(("1/8", "1/4", "1/2")):
        left = 1170 + index * 54
        add_rect(slide, left, 41, 45, 24, fill=source.NEUTRAL_FILL,
                 stroke=source.LIGHT_LINE, sw=1.0, radius=True,
                 name=f"decoder-scale-{label.replace('/', '-')}")
        add_text(slide, left + 22.5, 58, label, size=15,
                 color=source.MUTED, bold=True, width=40)
    add_line(slide, 1331, 53, 1360, 53, color=source.LINE,
             sw=1.4, arrow=True)
    add_text(slide, 1452, 58, "native route", size=15.5,
             color=source.MUTED, bold=True, width=130,
             name="native-route-label")

    add_voxel_cube(slide)
    add_text(slide, 405, 236, "current voxel", size=17, bold=True,
             width=130, name="current-voxel-label")
    add_text(slide, 405, 257, "feature  Qₜˢ", size=17,
             color=source.MUTED, width=130, name="current-feature-label")
    add_line(slide, 455, 171, 480, 171, sw=2.3, arrow=True)

    # DCA.
    add_module_panel(slide, 480, 88, 255, 223, color=source.DCA,
                     fill="#F7FBFE", title="DCA", question="Where needed?",
                     tab_width=70, name="dca")
    add_picture(slide, 494, 119, 126, 132, traces["dca"],
                name="dca-trace")
    for index in range(3):
        left = 643 + index * 8
        top = 139 - index * 7
        add_rect(slide, left, top, 48, 34, fill="#E4EEF7",
                 stroke=source.DCA, sw=1.0, radius=True,
                 name=f"dca-image-query-{index + 1}")
        add_line(slide, left + 7, top + 10, left + 39, top + 10,
                 color="#9CBAD3", sw=1.0)
        add_line(slide, left + 7, top + 21, left + 31, top + 21,
                 color="#9CBAD3", sw=1.0)
    add_line(slide, 647, 202, 688, 202, color=source.DCA,
             sw=2.0, arrow=True)
    add_rect(slide, 660, 224, 59, 39, fill=source.DCA_FILL,
             stroke=source.DCA, sw=1.2, radius=True, name="dca-candidates")
    for row in range(3):
        for col in range(5):
            if (row + 2 * col) % 4 == 0:
                add_rect(slide, 667 + col * 9, 231 + row * 9, 6, 6,
                         fill=source.DCA, stroke=source.DCA, sw=0.2,
                         radius=True)
    add_text(slide, 607, 290, "candidate-guided re-query", size=18,
             color=source.DCA, bold=True, width=230,
             name="dca-summary")
    add_line(slide, 736, 171, 762, 171, sw=2.3, arrow=True)

    # VVE.
    add_module_panel(slide, 762, 88, 255, 223, color=source.VVE,
                     fill="#FFFBF5", title="VVE",
                     question="Where to read?", tab_width=70, name="vve")
    add_picture(slide, 776, 119, 126, 132, traces["vve"],
                name="vve-trace")
    grid_x, grid_y, cell = 930, 137, 15
    for row in range(5):
        for col in range(5):
            fill = "#FFFFFF"
            if (row, col) == (3, 1):
                fill = "#C7D4E3"
            if (row, col) == (1, 3):
                fill = source.TRANSPORT_FILL
            add_rect(slide, grid_x + col * cell, grid_y + row * cell,
                     cell, cell, fill=fill, stroke="#9AA8B8", sw=0.7,
                     name=f"vve-grid-{row}-{col}")
    add_line(slide, grid_x + 1.5 * cell, grid_y + 3.5 * cell,
             grid_x + 3.35 * cell, grid_y + 1.65 * cell,
             color=source.VVE, sw=2.4, arrow=True, name="vve-displacement")
    add_text(slide, 967, 229, "x − τv̂", size=18,
             color=source.VVE, bold=True, italic=True, width=85,
             name="vve-address")
    add_text(slide, 889, 290, "local current–history match", size=18,
             color=source.VVE, bold=True, width=235, name="vve-summary")
    add_line(slide, 1018, 171, 1044, 171, sw=2.3, arrow=True)

    # VDSF and supervised P/T/R routing.
    add_module_panel(slide, 1044, 88, 467, 223, color=source.VDSF,
                     fill="#F7FCFB", title="VDSF",
                     question="Which source?", tab_width=76, name="vdsf")
    routes = [
        ("P", "Persist", "H(x)", source.PERSIST, source.PERSIST_FILL,
         traces["persist"]),
        ("T", "Transport", "H(x − τv̂)", source.TRANSPORT,
         source.TRANSPORT_FILL, traces["transport"]),
        ("R", "Refresh", "Qₜ(x)", source.REFRESH, source.REFRESH_FILL,
         traces["refresh"]),
    ]
    for index, (abbr, label, formula, color, fill, trace) in enumerate(routes):
        left = 1058 + index * 111
        add_rect(slide, left, 117, 101, 159, fill=fill, stroke=color,
                 sw=1.25, radius=True, name=f"route-{label.lower()}")
        add_circle(slide, left + 15, 134, 10, fill=color,
                   name=f"route-{label.lower()}-badge")
        add_text(slide, left + 15, 139, abbr, size=13, color="#FFFFFF",
                 bold=True, width=18)
        add_text(slide, left + 30, 139, label, size=15.5, bold=True,
                 align="left", width=68,
                 name=f"route-{label.lower()}-label")
        add_picture(slide, left + 8, 149, 85, 80, trace,
                    name=f"{label.lower()}-trace")
        add_text(slide, left + 50.5, 249, formula, size=17, italic=True,
                 width=92, name=f"route-{label.lower()}-formula")
        add_text(slide, left + 50.5, 269, f"α{abbr}", size=17,
                 color=color, bold=True, width=80,
                 name=f"route-{label.lower()}-weight")

    fusion_x = 1409
    for index, color in enumerate(
            (source.PERSIST, source.TRANSPORT, source.REFRESH)):
        add_line(slide, 1391, 143 + index * 40, fusion_x - 18, 194,
                 color=color, sw=2.0, arrow=True,
                 name=f"route-to-fusion-{index + 1}")
    add_circle(slide, fusion_x, 194, 21, fill=source.VDSF,
               name="soft-route-sum")
    add_text(slide, fusion_x, 201, "Σ", size=25, color="#FFFFFF",
             bold=True, width=32)
    add_line(slide, 1432, 194, 1457, 194, sw=2.0, arrow=True)
    for row in range(4):
        for col in range(4):
            active = (row + col) % 3 != 0
            fill = source.VDSF if active else "#C8E5E1"
            add_rect(slide, 1461 + col * 10, 175 + row * 10, 7, 7,
                     fill=fill, stroke=fill, sw=0.2, radius=True,
                     name=f"sparse-fusion-{row}-{col}")
    add_multiline(slide, 1421, 248, ["soft route", "sparse fusion"],
                  size=17, color=source.VDSF, bold=True, width=150,
                  line_height=19, name="vdsf-fusion-label")
    add_text(slide, 1277, 290, "supervised P/T/R source routing", size=18,
             color=source.VDSF, bold=True, width=330, name="vdsf-summary")

    # Training paths use the same restrained dashed expansion grammar as the
    # GenComm-style reference, but all content is specific to RoadOcc.
    training_specs = [
        (607, 328, 142, "candidate targets", source.DCA),
        (889, 328, 142, "velocity targets", source.VVE),
        (1277, 328, 178, "P/T/R targets", source.VDSF),
    ]
    for index, (cx, cy, width, label, color) in enumerate(training_specs):
        add_rect(slide, cx - width / 2, cy - 13, width, 27,
                 fill="#FFFFFF", stroke=color, sw=1.0, radius=True,
                 dash=True, name=f"training-target-{index + 1}")
        add_text(slide, cx, cy + 5, label, size=15.5, color=color,
                 bold=True, width=width - 10)
        add_line(slide, cx, cy - 14, cx, 311, color=color, sw=1.2,
                 arrow=True, dash=True)
    add_text(slide, 351, 333, "training only", size=14.5,
             color=source.MUTED, bold=True, italic=True, align="left",
             width=110, name="training-only-label")

    # Temporal memory loop. P/T/R source semantics stay in the VDSF panel;
    # this lane is reserved for the recurrent state lifecycle.
    lane_y = 363
    add_rect(slide, 348, lane_y, 1168, 134, fill="#FAFBFC",
             stroke=source.LIGHT_LINE, sw=1.3, radius=True,
             name="temporal-memory-loop")
    add_text(slide, 369, lane_y + 25, "TEMPORAL MEMORY LOOP", size=18,
             color=source.MUTED, bold=True, align="left", width=280,
             name="temporal-memory-heading")
    for index in range(3):
        left = 382 + index * 12
        top = lane_y + 51 - index * 6
        add_rect(slide, left, top, 84, 41, fill="#EEF2F6",
                 stroke=source.NEUTRAL, sw=1.1, radius=True,
                 name=f"fifo-slot-{index + 1}")
        add_text(slide, left + 42, top + 26, f"slot  t−{3-index}",
                 size=16.5, bold=True, width=76)
    add_text(slide, 430, lane_y + 116, "cached feature + VVE prior",
             size=17, color=source.MUTED, bold=True, width=190,
             name="fifo-label")
    add_line(slide, 486, lane_y + 72, 518, lane_y + 72,
             sw=1.9, arrow=True)
    add_rect(slide, 524, lane_y + 45, 332, 55, fill="#FFFFFF",
             stroke=source.NEUTRAL, sw=1.2, radius=True,
             name="same-grid-memory")
    add_multiline(slide, 690, lane_y + 67,
                  ["same-grid multi-slot memory",
                   "fixed roadside coordinates"], size=17,
                  bold=True, width=310, line_height=19,
                  name="same-grid-memory-label")
    add_line(slide, 856, lane_y + 61, 882, lane_y + 61,
             color=source.VVE, sw=1.8, arrow=True)
    add_rect(slide, 888, lane_y + 43, 151, 35,
             fill=source.VVE_FILL, stroke=source.VVE, sw=1.0,
             radius=True, name="nearest-slot-vve")
    add_text(slide, 963.5, lane_y + 66, "nearest slot → VVE", size=16.5,
             color=source.VVE, bold=True, width=143)
    add_line(slide, 856, lane_y + 84, 882, lane_y + 84,
             color=source.VDSF, sw=1.8, arrow=True)
    add_rect(slide, 888, lane_y + 82, 174, 35,
             fill=source.VDSF_FILL, stroke=source.VDSF, sw=1.0,
             radius=True, name="all-slots-vdsf")
    add_text(slide, 975, lane_y + 105, "all valid slots → VDSF",
             size=16.5, color=source.VDSF, bold=True, width=166)

    add_elbow(slide, [(1005, 311), (1005, lane_y + 30),
                      (1195, lane_y + 30), (1195, lane_y + 47)],
              color=source.VVE, sw=1.7, arrow=True,
              name="vve-to-current-state")
    add_text(slide, 1090, lane_y + 23, "VVE prior  v̂ₜˢ", size=16.5,
             color=source.VVE, bold=True, width=160)
    add_rect(slide, 1195, lane_y + 45, 280, 58, fill=source.VDSF_FILL,
             stroke=source.VDSF, sw=1.2, radius=True,
             name="append-current-state")
    add_text(slide, 1335, lane_y + 66, "append current state", size=17,
             color=source.VDSF, bold=True, width=250)
    add_text(slide, 1335, lane_y + 89, "fused Hₜˢ  +  VVE v̂ₜˢ", size=16.5,
             bold=True, width=250, name="append-current-state-label")
    add_elbow(slide, [(1409, 312), (1409, lane_y + 44)],
              color=source.VDSF, sw=1.8, arrow=True,
              name="vdsf-to-current-state")

    add_elbow(slide, [(1440, lane_y + 104), (1440, lane_y + 122),
                      (430, lane_y + 122), (430, lane_y + 101)],
              color=source.NEUTRAL, sw=1.5, arrow=True,
              name="next-frame-return")
    add_rect(slide, 775, lane_y + 111, 320, 22, fill="#FAFBFC",
             stroke="#FAFBFC", sw=0.0, name="next-frame-label-bg")
    add_text(slide, 935, lane_y + 127,
             "reuse as history at  t+1", size=16.5,
             color=source.MUTED, bold=True, width=310,
             name="next-frame-label")

    # Prediction panel.
    add_line(slide, 1536, 171, 1562, 171, sw=2.3, arrow=True,
             name="decoder-to-output")
    add_text(slide, 1680, 32, "Predictions", size=24, bold=True,
             width=210, name="predictions-heading")
    add_rect(slide, 1565, 50, 219, 447, fill="#FFFFFF",
             stroke=source.LINE, sw=1.5, radius=True, name="predictions")
    add_picture(slide, 1578, 70, 193, 183, traces["occupancy"],
                name="occupancy-output")
    add_text(slide, 1674.5, 276, "semantic occupancy", size=18,
             bold=True, width=190, name="occupancy-output-label")
    for index, (x1, y1, x2, y2) in enumerate([
            (1605, 317, 1640, 304), (1650, 320, 1688, 320),
            (1695, 305, 1730, 319)]):
        add_line(slide, x1, y1, x2, y2, color=source.VVE,
                 sw=2.4, arrow=True, name=f"velocity-vector-{index + 1}")
    add_text(slide, 1674.5, 351, "voxel velocity", size=18,
             color=source.VVE, bold=True, width=180,
             name="velocity-output-label")
    add_rect(slide, 1592, 387, 165, 62, fill=source.NEUTRAL_FILL,
             stroke=source.LIGHT_LINE, sw=1.0, radius=True,
             name="full-resolution-state")
    add_multiline(slide, 1674.5, 411,
                  ["full-resolution", "scene state"], size=17,
                  color=source.MUTED, bold=True, width=150,
                  line_height=20, name="full-resolution-state-label")

    return slide


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--paper-dir", type=Path, default=Path("docs/paper/ICLR2027"),
        help="Directory containing the manuscript and figure assets.")
    parser.add_argument(
        "--output", type=Path, default=None,
        help="Output PPTX path. Defaults to figures/roadocc_framework_editable.pptx.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    paper_dir = args.paper_dir.resolve()
    output = args.output
    if output is None:
        output = paper_dir / "figures" / "roadocc_framework_editable.pptx"
    else:
        output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    prs.core_properties.title = "RoadOcc framework"
    prs.core_properties.subject = "Editable ICLR 2027 framework figure"
    prs.core_properties.comments = (
        "Native PowerPoint shapes with independent evidence image layers. "
        "Visual grammar informed by recent occupancy and collaborative "
        "perception framework figures; no reference artwork is copied.")
    traces = source.load_existing_traces(paper_dir)
    build_slide(prs, traces)
    prs.save(output)
    print(f"Editable PowerPoint: {output}")
    print(f"Slide size: {SLIDE_W_IN:.3f} x {SLIDE_H_IN:.3f} in")


if __name__ == "__main__":
    main()
